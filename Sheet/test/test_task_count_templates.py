import importlib.util
from pathlib import Path
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("builder", ROOT / "build_four_workbooks.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class TaskCountTemplatesTest(unittest.TestCase):
    def test_reads_shared_code_config_and_rejects_malformed_values(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Path(directory) / "Code.gs"
            config.write_text("const DM = {defaultTaskCount: 30, practiceTaskCounts: {1:30, 2:25, 3:40}};", encoding="utf8")
            counts = builder.load_task_counts(config)
            self.assertEqual([counts[i] for i in range(1, 5)], [30, 25, 40, 30])
            for mapping in ['2: 0', '2: -1', '2: 2.5', '2: "25"', '2: 201', '16: 25', '"02":25', '2:25,2:40']:
                config.write_text("const DM = {defaultTaskCount: 30, practiceTaskCounts: {" + mapping + "}};", encoding="utf8")
                with self.subTest(mapping=mapping), self.assertRaises(ValueError):
                    builder.load_task_counts(config)
        self.assertEqual(len(builder.load_task_counts()), 15)

    def test_builds_three_workbooks_with_30_25_40_and_shared_formulas(self):
        counts = {practice: 30 for practice in range(1, 16)}
        counts.update({2: 25, 3: 40})
        with tempfile.TemporaryDirectory() as directory:
            for index, name in enumerate(builder.PRACTITIONERS, 1):
                path = Path(directory) / f"group-{index}.xlsx"
                builder.make_plus(path, name, index, counts)
                workbook = load_workbook(path)
                first = 5 if name == 'Рами' else 6
                solved = 2 if name == 'Рами' else 3
                for practice, count in counts.items():
                    sheet = workbook[f"Практика {practice}"]
                    last = first + count - 1
                    self.assertEqual(sheet.max_column, last)
                    self.assertEqual([sheet.cell(3, c).value for c in range(first, last + 1)], [str(i) for i in range(1, count + 1)])
                    self.assertIn(f"{get_column_letter(first)}4:{get_column_letter(last)}4", sheet.cell(4, solved).value)
                    self.assertIn(f"{get_column_letter(first)}$3:{get_column_letter(last)}$3", sheet.cell(4, solved + 1).value)
                    rules = [rule for group in sheet.conditional_formatting for rule in sheet.conditional_formatting[group]]
                    self.assertTrue(any(rule.formula == [f"2*{count}/3"] for rule in rules))
                    self.assertTrue(any(rule.formula == [f"{count}/2", f"2*{count}/3"] for rule in rules))
                    self.assertFalse(sheet.protection.sheet)
                    self.assertIsNone(sheet.cell(4, last).value)
                self.assertIn('Логи', workbook.sheetnames)
                workbook.close()


if __name__ == '__main__':
    unittest.main()
